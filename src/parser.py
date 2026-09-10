import os
import re
from datetime import datetime
import pymupdf  # PyMuPDF
import pypdf
from .utils import (
    SKILLS_TAXONOMY, DEGREE_PATTERNS, clean_text,
    extract_email, extract_phone, extract_linkedin, extract_github
)

class ResumeParser:
    def __init__(self):
        self.skills_taxonomy = SKILLS_TAXONOMY

    def extract_text(self, file_path_or_stream, filename="document.pdf"):
        """Extract text from PDF, DOCX, or TXT file path or stream."""
        ext = os.path.splitext(filename)[1].lower() if filename else ".pdf"

        # 1. Plain text files
        if ext in [".txt", ".text"]:
            try:
                if isinstance(file_path_or_stream, str):
                    with open(file_path_or_stream, "r", encoding="utf-8", errors="ignore") as f:
                        return clean_text(f.read())
                else:
                    content = file_path_or_stream.read()
                    if isinstance(content, bytes):
                        content = content.decode("utf-8", errors="ignore")
                    file_path_or_stream.seek(0)
                    return clean_text(content)
            except Exception as e:
                print(f"Error reading TXT: {e}")

        # 2. Word documents (.docx)
        if ext in [".docx", ".doc"]:
            try:
                import docx
                doc = docx.Document(file_path_or_stream)
                full_text = []
                for para in doc.paragraphs:
                    if para.text.strip():
                        full_text.append(para.text)
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                        if row_text:
                            full_text.append(row_text)
                return clean_text("\n".join(full_text))
            except Exception as e:
                print(f"Error reading DOCX: {e}")

        # 3. PDF fallback extraction
        return self.extract_text_from_pdf(file_path_or_stream)

    def extract_text_from_pdf(self, pdf_path_or_stream):
        """Extract text from PDF file path or stream using PyMuPDF with pypdf fallback."""
        raw_text = ""
        try:
            if isinstance(pdf_path_or_stream, str):
                doc = pymupdf.open(pdf_path_or_stream)
            else:
                doc = pymupdf.open(stream=pdf_path_or_stream.read(), filetype="pdf")
                pdf_path_or_stream.seek(0)
            
            for page in doc:
                raw_text += page.get_text("text") + "\n"
            doc.close()
        except Exception as e:
            # Fallback to pypdf
            try:
                reader = pypdf.PdfReader(pdf_path_or_stream)
                for page in reader.pages:
                    raw_text += (page.extract_text() or "") + "\n"
            except Exception as ex:
                print(f"Error reading PDF: {ex}")
        
        return clean_text(raw_text)

    def parse(self, file_path_or_stream, filename="Resume.pdf"):
        """Main parsing method returning a structured candidate profile dict."""
        raw_text = self.extract_text(file_path_or_stream, filename=filename)
        lines = [line.strip() for line in raw_text.split('\n') if line.strip()]

        name = self._extract_name(lines, filename)
        title = self._extract_title(lines)
        contact_info = self._extract_contact(raw_text)
        skills = self._extract_skills(raw_text)
        education_info = self._extract_education(raw_text)
        exp_years, work_history = self._extract_experience(raw_text)

        cand_id = os.path.splitext(os.path.basename(filename))[0]
        # Deterministic short anonymized ID for blind screening
        import hashlib
        short_hash = hashlib.md5(cand_id.encode('utf-8')).hexdigest()[:4].upper()
        anonymous_id = f"CAND-{short_hash}"

        # Generate concise recruiter pitch
        top_skills_str = ", ".join(skills[:4]) if skills else "Core software engineering"
        summary_pitch = f"{title} with {exp_years} yrs experience ({top_skills_str}) and {education_info['highest_degree']} degree."

        return {
            "id": cand_id,
            "anonymous_id": anonymous_id,
            "filename": filename,
            "name": name,
            "title": title,
            "contact": contact_info,
            "skills": skills,
            "experience_years": exp_years,
            "highest_degree": education_info["highest_degree"],
            "education": education_info["degrees"],
            "work_history": work_history,
            "raw_text": raw_text,
            "summary_pitch": summary_pitch,
            "status": "New"
        }

    def _extract_name(self, lines, filename):
        """Extract candidate name from header lines or fallback to filename."""
        if not lines:
            return os.path.splitext(os.path.basename(filename))[0].replace('_', ' ').title()
        
        for line in lines[:3]:
            clean_line = re.sub(r'(?i)\b(resume|cv|curriculum vitae|page \d+)\b', '', line).strip()
            if clean_line and 2 <= len(clean_line.split()) <= 5 and not re.search(r'\d|@|\.com|\.io', clean_line):
                if not re.search(r'(?i)\b(summary|experience|skills|education|profile|overview)\b', clean_line):
                    return clean_line
        
        base = os.path.splitext(os.path.basename(filename))[0]
        base = re.sub(r'^resume_\d+_', '', base)
        return base.replace('_', ' ').title()

    def _extract_title(self, lines):
        """Extract candidate headline/title."""
        if len(lines) > 1:
            for line in lines[1:5]:
                if not re.search(r'@|\d{3}-|\.com|http|github|linkedin', line):
                    if len(line.split()) <= 8 and not re.search(r'(?i)\b(summary|overview|skills|experience|education)\b', line):
                        return line
        return "Software Professional"

    def _extract_contact(self, text):
        email = extract_email(text)
        phone = extract_phone(text)
        linkedin = extract_linkedin(text)
        github = extract_github(text)
        
        location = None
        loc_match = re.search(r'\b([A-Z][a-zA-Z\s]+,\s*[A-Z]{2})\b', text)
        if loc_match:
            location = loc_match.group(1).strip()

        return {
            "email": email or "N/A",
            "phone": phone or "N/A",
            "linkedin": linkedin or "N/A",
            "github": github or "N/A",
            "location": location or "Not specified"
        }

    def _extract_skills(self, text):
        """Match text against canonical skills taxonomy."""
        text_lower = text.lower()
        matched_skills = set()
        
        for keyword, canonical in self.skills_taxonomy.items():
            escaped = re.escape(keyword)
            pattern = r'(?i)(?:\b|_)' + escaped + r'(?:\b|_)'
            if keyword in ['c++', 'c#', '.net', 'next.js', 'node.js', 'vue.js']:
                pattern = r'(?i)' + escaped
            
            if re.search(pattern, text_lower):
                matched_skills.add(canonical)
                
        return sorted(list(matched_skills))

    def _extract_education(self, text):
        degrees = []
        highest_degree = "None"
        max_level = 0
        
        for pattern, degree_name, level in DEGREE_PATTERNS:
            matches = re.finditer(pattern, text)
            for m in matches:
                if level > max_level:
                    max_level = level
                    highest_degree = degree_name
                
                start = max(0, m.start() - 30)
                end = min(len(text), m.end() + 60)
                context = text[start:end].replace('\n', ' ').strip()
                if degree_name not in [d["degree"] for d in degrees]:
                    degrees.append({
                        "degree": degree_name,
                        "context": context
                    })
                    
        return {
            "highest_degree": highest_degree,
            "degrees": degrees
        }

    def _extract_experience(self, text):
        """Parse work history date ranges specifically within work experience sections."""
        current_year = datetime.now().year
        
        # 1. Prefer explicit mentioned years in summary like "5.5 years of experience" or "7+ years of experience"
        exp_mention_match = re.search(r'(?i)(\d+(?:\.\d+)?)\s*\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:experience|industry\s+experience)', text)
        stated_years = None
        if exp_mention_match:
            try:
                stated_years = float(exp_mention_match.group(1))
            except ValueError:
                pass

        # 2. Isolate work experience section text
        exp_section = text
        exp_heading_match = re.search(r'(?i)\b(work experience|professional experience|employment history|experience timeline|work history)\b', text)
        edu_heading_match = re.search(r'(?i)\b(education|academic background|education background|degrees)\b', text)

        if exp_heading_match:
            start_pos = exp_heading_match.end()
            end_pos = edu_heading_match.start() if (edu_heading_match and edu_heading_match.start() > start_pos) else len(text)
            exp_section = text[start_pos:end_pos]

        # Match date ranges like "Mar 2022 - Present", "2019 - 2022", "Jun 2020 - Jul 2022"
        date_range_pattern = r'(?i)\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)?\s*(\d{4})\s*[-\u2013\u2014to]+\s*(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)?\s*(\d{4}|Present|Current)\b'
        
        matches = re.findall(date_range_pattern, exp_section)
        total_months = 0
        
        month_map = {
            'jan': 1, 'feb': 2, 'mar': 3, 'apr': 4, 'may': 5, 'jun': 6,
            'jul': 7, 'aug': 8, 'sep': 9, 'oct': 10, 'nov': 11, 'dec': 12
        }

        for start_m, start_y, end_m, end_y in matches:
            try:
                start_year = int(start_y)
                start_month = month_map.get(start_m.lower(), 1) if start_m else 1
                
                if end_y.lower() in ['present', 'current']:
                    end_year = current_year
                    end_month = datetime.now().month
                else:
                    end_year = int(end_y)
                    end_month = month_map.get(end_m.lower(), 12) if end_m else 12
                
                months = (end_year - start_year) * 12 + (end_month - start_month)
                if 0 < months <= 240:  # Sanity cap per position
                    total_months += months
            except ValueError:
                continue

        calc_years = round(total_months / 12.0, 1)
        
        if stated_years is not None:
            final_years = stated_years
        elif calc_years > 0:
            final_years = calc_years
        else:
            final_years = 1.0

        # Extract job roles
        history = []
        lines = exp_section.split('\n')
        for line in lines:
            if re.search(r'(?i)\b(engineer|developer|architect|analyst|manager|lead|specialist|consultant)\b', line):
                if len(line.split()) <= 10 and line.strip() not in history:
                    history.append(line.strip())

        return final_years, history[:5]
